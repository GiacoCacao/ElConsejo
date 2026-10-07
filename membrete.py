"""Acta con membrete en PDF (ReportLab) y Word (python-docx).

El membrete (nombre de la institución, lema, ciudad, papel y logotipo opcional) se configura en
Ajustes. Sin logotipo se usa el sello de El Consejo: en el PDF dibujado en vectorial; en el Word, la
imagen static/sello.png. El texto del acta llega en el Markdown sencillo que compone sesiones.py.
"""
import io
import os
import re
from datetime import datetime
from zoneinfo import ZoneInfo

import bd
import config

AQUI = os.path.dirname(os.path.abspath(__file__))
TTF = os.path.join(AQUI, "static", "fuentes", "ttf")
SELLO_PNG = os.path.join(AQUI, "static", "sello.png")
DIR_MEMBRETE = os.path.join(config.DATOS, "membrete")
ORO, ORO_OSC, TINTA, GRIS = "#b0894a", "#8a6a32", "#1f1b14", "#6b6250"
MESES = "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split()


def ajustes():
    return {"nombre": bd.ajuste("membrete_nombre", "El Consejo"),
            "lema": bd.ajuste("membrete_lema", "Consejo de expertos"),
            "ciudad": bd.ajuste("membrete_ciudad", "Caracas"),
            "papel": bd.ajuste("membrete_papel", "carta"),
            "logo": bool(logo())}


def logo():
    for ext in ("png", "jpg", "jpeg"):
        p = os.path.join(DIR_MEMBRETE, f"logo.{ext}")
        if os.path.exists(p):
            return p
    return None


def guardar_logo(nombre, datos):
    ext = os.path.splitext(nombre or "")[1].lower().lstrip(".")
    if ext not in ("png", "jpg", "jpeg"):
        raise ValueError("El logotipo debe ser PNG o JPG")
    if len(datos) > 3 * 1024 * 1024:
        raise ValueError("El logotipo supera 3 MB")
    borrar_logo()
    os.makedirs(DIR_MEMBRETE, exist_ok=True)
    with open(os.path.join(DIR_MEMBRETE, f"logo.{ext}"), "wb") as f:
        f.write(datos)


def borrar_logo():
    p = logo()
    if p:
        os.remove(p)


def _fecha_larga(ts):
    d = datetime.fromtimestamp(ts, ZoneInfo(config.TZ))
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def _bloques(md):
    """Markdown del acta → [(tipo, texto)] con tipo h1/h2/h3/p/ul/ol/cita."""
    out = []
    for linea in md.splitlines():
        l = linea.rstrip()
        if not l.strip() or l.strip() == "*La Secretaría del Consejo*":
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", l):
            out.append((f"h{len(m.group(1))}", m.group(2)))
        elif m := re.match(r"^>\s?(.*)", l):
            out.append(("cita", m.group(1)))
        elif m := re.match(r"^\s*[-*]\s+(.*)", l):
            out.append(("ul", m.group(1)))
        elif m := re.match(r"^\s*(\d+)\.\s+(.*)", l):
            out.append(("ol", m.group(2)))
        else:
            out.append(("p", l))
    return out


# ---- PDF ---------------------------------------------------------------------------
def _fuentes_pdf():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "CorSemi" in pdfmetrics.getRegisteredFontNames():
        return
    for alias, archivo in (("CorMed", "Cormorant-Medium"), ("CorSemi", "Cormorant-SemiBold"),
                           ("CorIt", "Cormorant-Italic"), ("Inter", "Inter-Regular"), ("InterSemi", "Inter-SemiBold")):
        pdfmetrics.registerFont(TTFont(alias, os.path.join(TTF, f"{archivo}.ttf")))
    pdfmetrics.registerFontFamily("Inter", normal="Inter", bold="InterSemi", italic="Inter", boldItalic="InterSemi")


def _en_linea_pdf(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    return re.sub(r"\*(.+?)\*", r'<font name="CorIt" size="+1.5">\1</font>', t)


def pdf(sesion, panel_nombre, acta_md):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)
    from reportlab.pdfgen import canvas as rl_canvas

    _fuentes_pdf()
    m = ajustes()
    papel = A4 if m["papel"] == "a4" else LETTER
    ancho, alto = papel
    ruta_logo = logo()
    num = sesion["numero"]

    est = {
        "titulo": ParagraphStyle("t", fontName="CorSemi", fontSize=25, leading=29, alignment=TA_CENTER, textColor=TINTA,
                                spaceAfter=7),
        "sub": ParagraphStyle("s", fontName="Inter", fontSize=8.8, leading=13, alignment=TA_CENTER, textColor=GRIS, spaceAfter=10),
        "h2": ParagraphStyle("h2", fontName="CorSemi", fontSize=14.5, leading=18, textColor=TINTA, spaceBefore=12, spaceAfter=2),
        "h3": ParagraphStyle("h3", fontName="CorMed", fontSize=12.5, leading=16, textColor=ORO_OSC, spaceBefore=8, spaceAfter=2),
        "p": ParagraphStyle("p", fontName="Inter", fontSize=9.3, leading=14.2, textColor=TINTA, spaceAfter=5),
        "li": ParagraphStyle("li", fontName="Inter", fontSize=9.3, leading=13.6, textColor=TINTA),
        "cita": ParagraphStyle("c", fontName="CorIt", fontSize=12, leading=16, textColor=TINTA),
        "firma": ParagraphStyle("f", fontName="Inter", fontSize=8, leading=11, alignment=TA_CENTER, textColor=GRIS),
    }

    def membrete(c, doc):
        c.saveState()
        x0, y0 = 2.2 * cm, alto - 1.55 * cm
        if ruta_logo:
            c.drawImage(ruta_logo, x0, y0 - 1.0 * cm, width=1.2 * cm, height=1.2 * cm, preserveAspectRatio=True, mask="auto")
        else:   # el sello de El Consejo, en vectorial
            cx, cy, r = x0 + 0.6 * cm, y0 - 0.4 * cm, 0.6 * cm
            c.setStrokeColor(colors.HexColor(ORO)); c.setLineWidth(0.8); c.circle(cx, cy, r)
            c.setLineWidth(0.4); c.circle(cx, cy, r * 0.82)
            c.arc(cx - r * 0.58, cy - r * 0.81, cx + r * 0.58, cy + r * 0.35, 0, 180)
            c.setFillColor(colors.HexColor(ORO_OSC)); c.setFont("CorSemi", 15); c.drawCentredString(cx, cy - 0.17 * cm, "C")
        c.setFillColor(colors.HexColor(TINTA)); c.setFont("CorSemi", 15)
        c.drawString(x0 + 1.55 * cm, y0 - 0.25 * cm, m["nombre"].upper(), charSpace=2.6)
        c.setFillColor(colors.HexColor(ORO_OSC)); c.setFont("Inter", 6.3)
        c.drawString(x0 + 1.58 * cm, y0 - 0.72 * cm, m["lema"].upper(), charSpace=1.8)
        c.setFillColor(colors.HexColor(GRIS)); c.setFont("Inter", 7)
        c.drawRightString(ancho - 2.2 * cm, y0 - 0.25 * cm, f"ACTA DE LA SESIÓN Nº {num}", charSpace=1.2)
        c.setFont("CorIt", 10); c.drawRightString(ancho - 2.2 * cm, y0 - 0.68 * cm, panel_nombre)
        c.setStrokeColor(colors.HexColor(ORO)); c.setLineWidth(0.7)
        c.line(2.2 * cm, y0 - 1.2 * cm, ancho - 2.2 * cm, y0 - 1.2 * cm)
        c.setLineWidth(0.25); c.line(2.2 * cm, y0 - 1.28 * cm, ancho - 2.2 * cm, y0 - 1.28 * cm)
        # pie
        c.setStrokeColor(colors.HexColor("#d9ccb0")); c.setLineWidth(0.4)
        c.line(2.2 * cm, 1.55 * cm, ancho - 2.2 * cm, 1.55 * cm)
        c.setFillColor(colors.HexColor(GRIS)); c.setFont("Inter", 7)
        c.drawString(2.2 * cm, 1.1 * cm, f"{m['nombre']} · {panel_nombre} · Sesión nº {num}")
        c.restoreState()

    class Paginado(rl_canvas.Canvas):   # «Página x de y»
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.setFont("Inter", 7); self.setFillColor(colors.HexColor(GRIS))
                self.drawRightString(ancho - 2.2 * cm, 1.1 * cm, f"Página {self._pageNumber} de {total}")
                super().showPage()
            super().save()

    historia = []
    lista, tipo_lista = [], None

    def cerrar_lista():
        nonlocal lista, tipo_lista
        if lista:
            historia.append(ListFlowable([ListItem(x, leftIndent=12) for x in lista],
                                         bulletType="bullet" if tipo_lista == "ul" else "1", start="–" if tipo_lista == "ul" else None,
                                         bulletFontName="Inter", bulletFontSize=8.5, bulletColor=colors.HexColor(ORO_OSC),
                                         leftIndent=14, spaceBefore=1, spaceAfter=6))
        lista, tipo_lista = [], None

    primero_p = True
    for tipo, texto in _bloques(acta_md):
        if tipo in ("ul", "ol"):
            if tipo_lista and tipo_lista != tipo:
                cerrar_lista()
            tipo_lista = tipo
            lista.append(Paragraph(_en_linea_pdf(texto), est["li"]))
            continue
        cerrar_lista()
        if tipo == "h1":
            historia += [Spacer(1, 4), Paragraph(_en_linea_pdf(texto), est["titulo"])]
        elif tipo == "h2":
            historia += [Paragraph(_en_linea_pdf(texto), est["h2"]),
                         HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#d9ccb0"), spaceAfter=5)]
        elif tipo == "h3":
            historia.append(Paragraph(_en_linea_pdf(texto), est["h3"]))
        elif tipo == "cita":
            t = Table([[Paragraph(_en_linea_pdf(texto), est["cita"])]], colWidths=["100%"])
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f6f0e3")),
                                   ("LINEBEFORE", (0, 0), (0, -1), 1.6, colors.HexColor(ORO)),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                                   ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            historia += [t, Spacer(1, 6)]
        else:
            historia.append(Paragraph(_en_linea_pdf(texto), est["sub"] if primero_p else est["p"]))
            primero_p = False
    cerrar_lista()

    # lugar, fecha y firmas
    fecha = _fecha_larga(sesion.get("cerrada") or sesion.get("abierta"))
    firma = lambda rol: [Spacer(1, 30), HRFlowable(width="80%", thickness=0.5, color=colors.HexColor(GRIS)),
                         Paragraph(rol, est["firma"])]
    t = Table([[firma("La Presidencia"), firma("La Secretaría del Consejo")]], colWidths=["50%", "50%"])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    historia.append(KeepTogether([Spacer(1, 14), Paragraph(f"En {m['ciudad']}, a {fecha}.", est["p"]), t]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=papel, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=3.4 * cm,
                            bottomMargin=2.3 * cm, title=f"Acta de la sesión nº {num} · {panel_nombre}",
                            author=m["nombre"], subject=sesion.get("asunto") or "")
    doc.build(historia, onFirstPage=membrete, onLaterPages=membrete, canvasmaker=Paginado)
    return buf.getvalue()


# ---- Word -----------------------------------------------------------------------------
def _runs_docx(par, texto, tam=None, fuente="Inter"):
    from docx.shared import Pt
    for trozo in re.split(r"(\*\*.+?\*\*|\*.+?\*)", texto):
        if not trozo:
            continue
        negrita, cursiva = trozo.startswith("**"), trozo.startswith("*") and not trozo.startswith("**")
        r = par.add_run(trozo.strip("*") if negrita or cursiva else trozo)
        r.bold = negrita
        r.italic = cursiva
        r.font.name = "Cormorant Garamond" if cursiva else fuente
        if tam:
            r.font.size = Pt(tam + (1.5 if cursiva else 0))


def _campo(par, codigo):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for tipo, texto in (("begin", None), (None, codigo), ("separate", None), ("end", None)):
        r = par.add_run()
        if tipo:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = f" {texto} "
        r._r.append(e)


def _borde_inferior(par, color="B0894A", grosor="8"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", grosor), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b); pPr.append(bdr)


def docx(sesion, panel_nombre, acta_md):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    m = ajustes()
    d = Document()
    sec = d.sections[0]
    if m["papel"] == "a4":
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    else:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin, sec.header_distance = Cm(3.2), Cm(2.3), Cm(1.1)
    base = d.styles["Normal"]
    base.font.name, base.font.size = "Inter", Pt(10)
    base.font.color.rgb = RGBColor.from_string(TINTA[1:])

    # membrete
    cab = sec.header
    t = cab.add_table(rows=1, cols=2, width=sec.page_width - sec.left_margin - sec.right_margin)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    izq, der = t.rows[0].cells
    izq.width, der.width = Cm(11), Cm(6.2)
    p = izq.paragraphs[0]
    p.add_run().add_picture(logo() or SELLO_PNG, height=Cm(1.15))
    r = p.add_run("   " + m["nombre"].upper()); r.font.name = "Cormorant Garamond"; r.font.size = Pt(15); r.bold = True
    p2 = izq.add_paragraph(); r = p2.add_run(m["lema"].upper()); r.font.size = Pt(6.5)
    r.font.color.rgb = RGBColor.from_string(ORO_OSC[1:])
    q = der.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = q.add_run(f"ACTA DE LA SESIÓN Nº {sesion['numero']}"); r.font.size = Pt(7); r.font.color.rgb = RGBColor.from_string(GRIS[1:])
    q2 = der.add_paragraph(); q2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = q2.add_run(panel_nombre); r.italic = True; r.font.name = "Cormorant Garamond"; r.font.size = Pt(11)
    linea = cab.add_paragraph(); _borde_inferior(linea)
    # pie con paginación
    pie = sec.footer.paragraphs[0]
    r = pie.add_run(f"{m['nombre']} · {panel_nombre} · Sesión nº {sesion['numero']}        Página ")
    r.font.size = Pt(7.5)
    _campo(pie, "PAGE")
    pie.add_run(" de ").font.size = Pt(7.5)
    _campo(pie, "NUMPAGES")

    primero_p = True
    for tipo, texto in _bloques(acta_md):
        if tipo == "h1":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(texto); r.font.name = "Cormorant Garamond"; r.font.size = Pt(24); r.bold = True
        elif tipo in ("h2", "h3"):
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(12 if tipo == "h2" else 8)
            r = p.add_run(texto); r.font.name = "Cormorant Garamond"; r.bold = tipo == "h2"
            r.font.size = Pt(14.5 if tipo == "h2" else 12.5)
            if tipo == "h2":
                _borde_inferior(p, "D9CCB0", "4")
            else:
                r.font.color.rgb = RGBColor.from_string(ORO_OSC[1:])
        elif tipo == "cita":
            p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            _runs_docx(p, f"*{texto}*" if not texto.startswith("*") else texto, 11)
        elif tipo in ("ul", "ol"):
            p = d.add_paragraph(style="List Bullet" if tipo == "ul" else "List Number")
            _runs_docx(p, texto, 10)
        else:
            p = d.add_paragraph()
            if primero_p:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER; primero_p = False
            _runs_docx(p, texto, 10)

    fecha = _fecha_larga(sesion.get("cerrada") or sesion.get("abierta"))
    d.add_paragraph().add_run(f"En {m['ciudad']}, a {fecha}.")
    f = d.add_table(rows=2, cols=2)
    for i, rol in enumerate(("La Presidencia", "La Secretaría del Consejo")):
        c1 = f.rows[0].cells[i].paragraphs[0]; c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c1.add_run("\n\n\n______________________________")
        c2 = f.rows[1].cells[i].paragraphs[0]; c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c2.add_run(rol); r.font.size = Pt(8.5)
    d.core_properties.title = f"Acta de la sesión nº {sesion['numero']} · {panel_nombre}"
    d.core_properties.author = m["nombre"]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
